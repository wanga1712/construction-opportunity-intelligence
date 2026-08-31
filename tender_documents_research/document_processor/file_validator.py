from pathlib import Path
from utils.logger_config import get_logger

def validate_open(path: Path, logger=None) -> bool:
    """
    Проверяет, что файл валиден и может быть открыт соответствующим парсером.
    Поддерживаются форматы: .pdf, .docx, .xlsx, .xlsm.
    Остальные форматы считаются валидными по умолчанию.
    """
    logger = logger or get_logger()
    ext = path.suffix.lower()

    try:
        if ext == ".pdf":
            from PyPDF2 import PdfReader
            PdfReader(str(path))
        elif ext == ".docx":
            import docx
            docx.Document(str(path))
        elif ext in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            load_workbook(filename=str(path), read_only=True, data_only=True)
        return True
    except Exception as e:
        logger.error(f"Файл поврежден ({ext}): {path.name} - {str(e)}")
        return False
