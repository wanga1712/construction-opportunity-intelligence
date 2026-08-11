
from pathlib import Path
from contextlib import contextmanager
from types import SimpleNamespace

import pytest

from document_processor.daemon import DocumentProcessorDaemon
from document_processor.downloader import Downloader
from document_processor.dto import ProcessingOutcome, TaskProcessResult


class _Repo:
    def __init__(self, status=None, local_path=None):
        self.status = status
        self.local_path = local_path
        self.marked = []
        self.attempts = []
        self.ensure = []
        self.finalize = []

    def ensure_download_file(self, queue_id, procurement_id, table_source, url, url_hash, file_name, source_id=None):
        self.ensure.append((queue_id, procurement_id, source_id, url_hash, file_name))

    def get_file_status(self, procurement_id, table_source, file_name, url_hash):
        if self.status:
            return (self.status, self.local_path)
        return None

    def mark_file_status(self, procurement_id, table_source, file_name, url_hash, status, worker_id=None):
        self.marked.append((procurement_id, file_name, url_hash, status))
        self.status = status

    def finalize_download_status(self, procurement_id, table_source, file_name, url_hash, success, error_message=None, local_path=None):
        self.finalize.append((procurement_id, file_name, url_hash, success, error_message, local_path))
        self.status = "COMPLETED" if success else "FAILED"
        self.local_path = str(local_path) if local_path else self.local_path

    def record_download_attempt(self, *args, **kwargs):
        self.attempts.append((args, kwargs))


class _HttpClient:
    def __init__(self):
        self.gets = 0

    def sanitize_name(self, value):
        return value

    def predict_filename(self, url):
        return "source.zip"


def _downloader(tmp_path, storage_root, repo, download_bytes=b"new"):
    d = Downloader.__new__(Downloader)
    d.base_dir = storage_root / "downloads-open"
    d.base_dir.mkdir(parents=True, exist_ok=True)
    d.approved_storage_root = storage_root.resolve()
    d.state_repo = repo
    d.http_client = _HttpClient()
    d.archive_extractor = SimpleNamespace(is_archive=lambda p: False)
    d.logger = SimpleNamespace(debug=lambda *a, **k: None, info=lambda *a, **k: None, warning=lambda *a, **k: None, error=lambda *a, **k: None)
    d._is_rar_part = lambda path: False

    def _download_single(task_dir, url, suggested_filename=None):
        d.http_client.gets += 1
        if download_bytes is None:
            return None
        task_dir.mkdir(parents=True, exist_ok=True)
        path = task_dir / (suggested_filename or "source.zip")
        path.write_bytes(download_bytes)
        return path

    d._download_single = _download_single
    return d


def test_s13_single_task_path_enters_s13_pipeline_not_legacy_task_kwargs():
    daemon = DocumentProcessorDaemon.__new__(DocumentProcessorDaemon)
    daemon.logger = SimpleNamespace(info=lambda *a, **k: None, warning=lambda *a, **k: None, error=lambda *a, **k: None)
    daemon.match_engine = object()
    calls = {"prefetch": 0, "s13": 0, "legacy": 0, "persist": 0, "completed": 0}

    daemon.pipeline = SimpleNamespace(
        prefetch_task=lambda *args, **kwargs: SimpleNamespace(files=[Path("fixture.txt")], failed_count=0, failures=[]),
        process_task_with_files=lambda *args, **kwargs: calls.__setitem__("legacy", calls["legacy"] + 1),
    )

    def process_task(**kwargs):
        calls["s13"] += 1
        assert kwargs["queue_id"] == 1
        assert kwargs["procurement_id"] == 1282
        assert kwargs["contract_reg_number"] == "0173200001424001779"
        assert kwargs["table_source"] == "reestr_contract_44_fz_awarded"
        assert kwargs["files"] == [Path("fixture.txt")]
        return TaskProcessResult(procurement_id=1282, queue_id=1, outcome=ProcessingOutcome.SUCCESS, files=[SimpleNamespace(status="COMPLETED")])

    daemon.s13_pipeline = SimpleNamespace(process_task=process_task)
    daemon.s13_persistence = SimpleNamespace(persist_task_result=lambda result: calls.__setitem__("persist", calls["persist"] + 1))
    daemon.s13_backend = SimpleNamespace(queue=SimpleNamespace(mark_completed=lambda task_id: calls.__setitem__("completed", calls["completed"] + 1), mark_failed=lambda *a, **k: None, mark_no_links=lambda *a, **k: None, mark_pending=lambda *a, **k: None))

    daemon._process_s13v2_task({
        "id": 1,
        "procurement_id": 1282,
        "source_id": 316812,
        "source_table": "reestr_contract_44_fz_awarded",
        "contract_number": "0173200001424001779",
    })

    assert calls["s13"] == 1
    assert calls["legacy"] == 0
    assert calls["persist"] == 1
    assert calls["completed"] == 0  # persistence owns atomic completion guard


def test_s13_download_root_uses_storage_root_and_invalid_root_fails_closed(tmp_path, monkeypatch):
    daemon = DocumentProcessorDaemon.__new__(DocumentProcessorDaemon)
    root = tmp_path / "data"
    root.mkdir()
    monkeypatch.setenv("DOCUMENT_STORAGE_ROOT", str(root))
    monkeypatch.delenv("DOCUMENT_DOWNLOAD_DIR", raising=False)
    assert daemon._resolve_download_base_dir("S13_V2") == root / "downloads-open"

    monkeypatch.setenv("DOCUMENT_DOWNLOAD_DIR", str(tmp_path / "elsewhere"))
    with pytest.raises(RuntimeError):
        daemon._resolve_download_base_dir("S13_V2")

    monkeypatch.delenv("DOCUMENT_STORAGE_ROOT", raising=False)
    monkeypatch.delenv("DOCUMENT_DOWNLOAD_DIR", raising=False)
    with pytest.raises(RuntimeError):
        daemon._resolve_download_base_dir("S13_V2")


def test_s13_off_root_completed_row_is_not_reused_and_redownloads_to_data(tmp_path):
    data_root = tmp_path / "data"
    off_root = tmp_path / "opt" / "downloads" / "1"
    off_root.mkdir(parents=True)
    off_file = off_root / "source.zip"
    off_file.write_bytes(b"old")
    repo = _Repo(status="COMPLETED", local_path=str(off_file))
    d = _downloader(tmp_path, data_root, repo, b"new")

    files, failure = d._process_single_link(1, data_root / "downloads-open" / "1", "https://example.invalid/source.zip", "source.zip", 1282, "source_table", None, None, source_id=316812)

    assert failure is None
    assert d.http_client.gets == 1
    assert repo.marked[-1][-1] == "PENDING"
    assert files[0].resolve().is_relative_to(data_root.resolve())
    assert repo.finalize[-1][5].resolve().is_relative_to(data_root.resolve())


def test_s13_valid_data_completed_row_reuse_get_0(tmp_path):
    data_root = tmp_path / "data"
    data_file = data_root / "downloads-open" / "1" / "source.zip"
    data_file.parent.mkdir(parents=True)
    data_file.write_bytes(b"cached")
    repo = _Repo(status="COMPLETED", local_path=str(data_file))
    d = _downloader(tmp_path, data_root, repo, b"new")

    files, failure = d._process_single_link(1, data_file.parent, "https://example.invalid/source.zip", "source.zip", 1282, "source_table", None, None, source_id=316812)

    assert failure is None
    assert files == [data_file]
    assert d.http_client.gets == 0
    assert repo.attempts[0][0][5] == "SKIPPED"


def test_s13_durable_download_survives_parser_crash_and_retry_reuses_data_file(tmp_path, monkeypatch):
    """S13 regression: durable download state is committed before parser/result failure.

    Contract proven here:
    HTTP success -> document_files durable -> download_attempt durable -> controlled
    process/parser-boundary crash -> queue not COMPLETED -> durable state remains ->
    retry of the same procurement/url reuses valid approved-root file with HTTP GET=0.
    """
    import psycopg2

    from document_processor.task_pipeline import TaskPipeline
    from document_processor.backends.queue_repository import S13V2QueueRepository
    from document_processor.backends.state_repository import S13V2StateRepository
    from document_processor.backends.s13_persistence import S13V2TaskPersistenceService

    
    import os

    conn = psycopg2.connect(
        host=os.environ.get("DB_HOST_TENDER", "127.0.0.1"),
        port=os.environ.get("DB_PORT_TENDER", "5432"),
        dbname=os.environ.get("DB_DATABASE_TENDER", "document_intelligence"),
        user=os.environ.get("DB_USER_TENDER", "doc_worker"),
        password=os.environ.get("DB_PASSWORD_TENDER"),
    )
    conn.autocommit = False
    with conn.cursor() as cur:
        cur.execute(
            """
            CREATE TEMPORARY TABLE document_processing_queue (
                id INT PRIMARY KEY,
                procurement_id INT NOT NULL,
                source_table TEXT,
                source_id INT,
                contract_number TEXT,
                status TEXT,
                pipeline_generation TEXT,
                worker_id INT,
                started_at TIMESTAMP,
                completed_at TIMESTAMP,
                last_error TEXT,
                queue_lane TEXT,
                priority_score INT,
                research_action TEXT,
                research_depth TEXT,
                category_codes TEXT[]
            ) ON COMMIT PRESERVE ROWS;

            CREATE TEMPORARY TABLE document_files (
                id SERIAL PRIMARY KEY,
                queue_id INT REFERENCES document_processing_queue(id),
                procurement_id INT NOT NULL,
                source_table TEXT,
                source_id INT,
                url TEXT,
                url_hash TEXT NOT NULL,
                file_name TEXT,
                download_status TEXT,
                worker_id INT,
                error_message TEXT,
                local_path TEXT,
                file_size_bytes BIGINT,
                downloaded_at TIMESTAMP,
                pipeline_generation TEXT NOT NULL,
                UNIQUE(url_hash, pipeline_generation)
            ) ON COMMIT PRESERVE ROWS;

            CREATE TEMPORARY TABLE download_attempts (
                id SERIAL PRIMARY KEY,
                queue_id INT,
                procurement_id INT,
                file_id INT,
                source_url TEXT,
                url_hash TEXT,
                attempt_number INT,
                started_at TIMESTAMP,
                finished_at TIMESTAMP,
                http_status INT,
                error_class TEXT,
                bytes BIGINT,
                latency DOUBLE PRECISION,
                result TEXT,
                duration_ms INT,
                bytes_received BIGINT,
                pipeline_generation TEXT
            ) ON COMMIT PRESERVE ROWS;

            CREATE TEMPORARY TABLE document_processing_results (
                id SERIAL PRIMARY KEY,
                queue_id INT REFERENCES document_processing_queue(id),
                procurement_id INT,
                file_id INT REFERENCES document_files(id),
                status TEXT,
                pages_processed INT,
                sheets_processed INT,
                rows_extracted INT,
                matches_found INT,
                processed_file_name TEXT,
                processed_local_path TEXT,
                archive_member_path TEXT,
                is_archive_member BOOLEAN DEFAULT FALSE,
                pipeline_generation TEXT
            ) ON COMMIT PRESERVE ROWS;

            CREATE TEMPORARY TABLE document_matches (
                id SERIAL PRIMARY KEY,
                queue_id INT REFERENCES document_processing_queue(id),
                procurement_id INT,
                file_id INT REFERENCES document_files(id),
                result_id INT REFERENCES document_processing_results(id),
                document_name TEXT,
                archive_member_path TEXT,
                match_count INT,
                score FLOAT,
                pipeline_generation TEXT
            ) ON COMMIT PRESERVE ROWS;

            CREATE TEMPORARY TABLE document_match_details (
                id SERIAL PRIMARY KEY,
                match_id INT REFERENCES document_matches(id),
                procurement_id INT,
                category_code TEXT,
                subcategory_code TEXT,
                matched_term TEXT,
                term_type TEXT,
                score FLOAT,
                row_data JSONB,
                page_or_sheet TEXT,
                row_number INT,
                context_before JSONB,
                context_after JSONB,
                pipeline_generation TEXT
            ) ON COMMIT PRESERVE ROWS;

            CREATE TEMPORARY TABLE document_evidence (
                id SERIAL PRIMARY KEY,
                procurement_id INT,
                queue_id INT REFERENCES document_processing_queue(id),
                match_id INT REFERENCES document_matches(id),
                category_code TEXT,
                evidence_score FLOAT,
                match_count INT,
                next_stage TEXT,
                updated_at TIMESTAMP,
                pipeline_generation TEXT,
                UNIQUE(procurement_id, category_code, pipeline_generation)
            ) ON COMMIT PRESERVE ROWS;
            """
        )
        cur.execute(
            """
            INSERT INTO document_processing_queue
                (id, procurement_id, source_id, source_table, contract_number, status, pipeline_generation)
            VALUES
                (101, 202001, 303001, 'reestr_contract_44_fz_awarded', 'TEST_CONTRACT_001', 'PROCESSING', 'S13_V2')
            """
        )
    conn.commit()

    class _Pool:
        def __init__(self, shared_conn):
            self.shared_conn = shared_conn

        @contextmanager
        def get_cursor(self, db_name=None):
            cursor = self.shared_conn.cursor()
            try:
                yield cursor
                self.shared_conn.commit()
            except Exception:
                self.shared_conn.rollback()
                raise
            finally:
                cursor.close()

    storage_root = tmp_path / "data" / "tender-documents"
    storage_root.mkdir(parents=True)
    monkeypatch.setenv("PROCESSING_BACKEND", "S13_V2")
    monkeypatch.setenv("DOCUMENT_STORAGE_ROOT", str(storage_root))
    monkeypatch.delenv("DOCUMENT_DOWNLOAD_DIR", raising=False)
    monkeypatch.delenv("REPROCESS_COMPLETED", raising=False)
    monkeypatch.setattr("document_processor.file_validator.validate_open", lambda path, logger: True)

    state_repo = S13V2StateRepository({}, pipeline_generation="S13_V2")
    state_repo._conn = conn
    queue_repo = S13V2QueueRepository({})
    queue_repo._conn = conn
    pool = _Pool(conn)

    http_gets = {"count": 0}

    class _HttpClient:
        def sanitize_name(self, value):
            return value

        def predict_filename(self, url):
            return "fixture.txt"

    downloader = Downloader.__new__(Downloader)
    downloader.base_dir = storage_root / "downloads-open"
    downloader.base_dir.mkdir(parents=True)
    downloader.approved_storage_root = storage_root.resolve()
    downloader.state_repo = state_repo
    downloader.http_client = _HttpClient()
    downloader.archive_extractor = SimpleNamespace(is_archive=lambda path: False)
    downloader.yandex_client = SimpleNamespace(build_remote_dir_and_prefix=lambda *a, **k: (None, None))
    downloader.contract_locator = SimpleNamespace(resolve_tender_id=lambda *a, **k: 303001)
    downloader.logger = SimpleNamespace(debug=lambda *a, **k: None, info=lambda *a, **k: None, warning=lambda *a, **k: None, error=lambda *a, **k: None)
    downloader._is_rar_part = lambda path: False

    def _download_single(task_dir, url, suggested_filename=None):
        http_gets["count"] += 1
        task_dir.mkdir(parents=True, exist_ok=True)
        path = task_dir / (suggested_filename or "fixture.txt")
        path.write_text("supported parser fixture", encoding="utf-8")
        return path

    downloader._download_single = _download_single
    downloader.get_links = lambda contract_number, table_source: [("https://example.invalid/s13/fixture.txt", "fixture.txt")]
    downloader.build_registry_link = lambda table_source, contract_number: "https://example.invalid/card"

    daemon = DocumentProcessorDaemon.__new__(DocumentProcessorDaemon)
    daemon.logger = SimpleNamespace(info=lambda *a, **k: None, warning=lambda *a, **k: None, error=lambda *a, **k: None)
    daemon.match_engine = object()
    daemon.pipeline = TaskPipeline(
        db=None,
        downloader=downloader,
        parser_factory=None,
        matcher=None,
        enhancer=None,
        worker_id=16,
        failed_uploads_dir=tmp_path / "failed",
        logger=daemon.logger,
        is_over_memory_limit=lambda: False,
    )
    daemon.s13_backend = SimpleNamespace(queue=queue_repo)
    daemon.s13_persistence = S13V2TaskPersistenceService(pool)

    parser_calls = {"count": 0}

    def _crashing_process_task(**kwargs):
        parser_calls["count"] += 1
        assert kwargs["queue_id"] == 101
        assert kwargs["procurement_id"] == 202001
        assert kwargs["files"]
        raise RuntimeError("TEST_PARSER_CRASH")

    daemon.s13_pipeline = SimpleNamespace(process_task=_crashing_process_task)
    task = {
        "id": 101,
        "procurement_id": 202001,
        "source_id": 303001,
        "source_table": "reestr_contract_44_fz_awarded",
        "contract_number": "TEST_CONTRACT_001",
    }

    daemon._process_s13v2_task(task)

    assert parser_calls["count"] == 1
    assert http_gets["count"] == 1

    with conn.cursor() as cur:
        cur.execute("SELECT status, last_error FROM document_processing_queue WHERE id=101")
        status_after_crash, last_error = cur.fetchone()
        assert status_after_crash == "FAILED"
        assert "TEST_PARSER_CRASH" in last_error
        cur.execute("SELECT id, procurement_id, source_id, download_status, local_path, pipeline_generation FROM document_files")
        file_id, procurement_id, source_id, download_status, local_path, pipeline_generation = cur.fetchone()
        cur.execute("SELECT result, attempt_number FROM download_attempts ORDER BY id")
        attempts_after_crash = cur.fetchall()
        cur.execute("SELECT COUNT(*) FROM document_processing_results")
        results_after_crash = cur.fetchone()[0]
        cur.execute("SELECT COUNT(*) FROM document_matches")
        matches_after_crash = cur.fetchone()[0]
        cur.execute("SELECT COUNT(*) FROM document_match_details")
        details_after_crash = cur.fetchone()[0]
        cur.execute("SELECT COUNT(*) FROM document_evidence")
        evidence_after_crash = cur.fetchone()[0]

    durable_path = Path(local_path)
    assert file_id > 0
    assert procurement_id == 202001
    assert source_id == 303001
    assert download_status == "COMPLETED"
    assert pipeline_generation == "S13_V2"
    assert durable_path.exists()
    assert durable_path.resolve().is_relative_to(storage_root.resolve())
    assert "opt/tender_documents_research/downloads" not in str(durable_path)
    assert attempts_after_crash == [("SUCCESS", 1)]
    assert results_after_crash == 0
    assert matches_after_crash == 0
    assert details_after_crash == 0
    assert evidence_after_crash == 0

    with conn.cursor() as cur:
        cur.execute("UPDATE document_processing_queue SET status='PROCESSING', last_error=NULL WHERE id=101")
    conn.commit()

    def _successful_process_task(**kwargs):
        parser_calls["count"] += 1
        assert kwargs["files"] == [durable_path]
        return TaskProcessResult(
            procurement_id=202001,
            queue_id=101,
            outcome=ProcessingOutcome.SUCCESS,
            files=[SimpleNamespace(file_name="fixture.txt", status="COMPLETED", error_message=None, pages=1, sheets=0, rows=0, matches=[])],
            evidence=[],
        )

    daemon.s13_pipeline = SimpleNamespace(process_task=_successful_process_task)
    daemon._process_s13v2_task(task)

    assert parser_calls["count"] == 2
    assert http_gets["count"] == 1  # retry HTTP GET = 0

    with conn.cursor() as cur:
        cur.execute("SELECT status, last_error FROM document_processing_queue WHERE id=101")
        final_status, final_error = cur.fetchone()
        cur.execute("SELECT result, attempt_number FROM download_attempts ORDER BY id")
        all_attempts = cur.fetchall()
        cur.execute("SELECT COUNT(*) FROM document_processing_results WHERE queue_id=101")
        final_results = cur.fetchone()[0]

    assert final_status == "COMPLETED"
    assert final_error is None
    assert final_results == 1
    assert all_attempts == [("SUCCESS", 1), ("SKIPPED", 0)]



def test_s13_pipeline_zero_match_parse_is_successful_processed_document(tmp_path):
    from document_processor.parser_factory import ParserFactory
    from document_processor.pipelines.s13_v2_pipeline import S13V2Pipeline

    doc = tmp_path / "no_match.txt"
    doc.write_text("Этот документ успешно читается, но целевых материалов здесь нет.", encoding="utf-8")

    class _NoMatchEngine:
        def process_text(self, text, line_meta=None):
            assert "успешно читается" in text
            return []

    pipeline = S13V2Pipeline(
        parser_factory=ParserFactory(),
        downloader=None,
        logger=SimpleNamespace(info=lambda *a, **k: None, warning=lambda *a, **k: None, error=lambda *a, **k: None),
        is_over_memory_limit=lambda: False,
    )

    result = pipeline.process_task(
        queue_id=501,
        procurement_id=202501,
        contract_reg_number="ZERO_MATCH",
        table_source="reestr_contract_44_fz_awarded",
        files=[doc],
        match_engine=_NoMatchEngine(),
    )

    assert result.outcome.value == "SUCCESS"
    assert len(result.files) == 1
    assert result.files[0].status == "COMPLETED"
    assert result.files[0].matches == []



def test_s13_pipeline_docx_word_parser_success_records_completed_file(tmp_path):
    import docx
    from document_processor.parser_factory import ParserFactory
    from document_processor.pipelines.s13_v2_pipeline import S13V2Pipeline

    doc_path = tmp_path / "word_fixture.docx"
    document = docx.Document()
    document.add_paragraph("Техническое задание успешно прочитано WordParser.")
    document.save(str(doc_path))

    class _NoMatchEngine:
        def process_text(self, text, line_meta=None):
            assert "WordParser" in text
            assert isinstance(line_meta, dict)
            return []

    pipeline = S13V2Pipeline(
        parser_factory=ParserFactory(),
        downloader=None,
        logger=SimpleNamespace(info=lambda *a, **k: None, warning=lambda *a, **k: None, error=lambda *a, **k: None),
        is_over_memory_limit=lambda: False,
    )

    result = pipeline.process_task(
        queue_id=503,
        procurement_id=202503,
        contract_reg_number="DOCX_WORD",
        table_source="reestr_contract_44_fz_awarded",
        files=[doc_path],
        match_engine=_NoMatchEngine(),
    )

    assert result.outcome.value == "SUCCESS"
    assert result.files[0].file_name == "word_fixture.docx"
    assert result.files[0].status == "COMPLETED"
    assert result.files[0].matches == []

def test_s13_pipeline_match_positive_produces_real_category_detail(tmp_path):
    from document_processor.match_engine import MatchEngine
    from document_processor.parser_factory import ParserFactory
    from document_processor.pipelines.s13_v2_pipeline import S13V2Pipeline

    doc = tmp_path / "positive.txt"
    doc.write_text("Ведомость материалов: композитный водоотводный лоток 100 м.", encoding="utf-8")
    engine = MatchEngine(
        keywords=["композитный водоотводный лоток"],
        keyword_meta={"композитный водоотводный лоток": {"category_codes": ["composite_drainage"]}},
        min_score=75,
    )
    pipeline = S13V2Pipeline(
        parser_factory=ParserFactory(),
        downloader=None,
        logger=SimpleNamespace(info=lambda *a, **k: None, warning=lambda *a, **k: None, error=lambda *a, **k: None),
        is_over_memory_limit=lambda: False,
    )

    result = pipeline.process_task(
        queue_id=502,
        procurement_id=202502,
        contract_reg_number="MATCH_POSITIVE",
        table_source="reestr_contract_44_fz_awarded",
        files=[doc],
        match_engine=engine,
    )

    assert result.files[0].status == "COMPLETED"
    assert result.files[0].matches
    detail = result.files[0].matches[0].details[0]
    assert detail.category_code in {"composite_drainage", "composites"}
    assert detail.matched_term


def test_s13_archive_extraction_preserves_durable_source_zip(tmp_path):
    import zipfile
    from document_processor.archive_extractor import ArchiveExtractor

    source_zip = tmp_path / "source.zip"
    with zipfile.ZipFile(source_zip, "w") as zf:
        zf.writestr("child.txt", "parseable child")

    extractor = ArchiveExtractor(SimpleNamespace(info=lambda *a, **k: None, warning=lambda *a, **k: None, error=lambda *a, **k: None))
    extracted = extractor.extract_recursive(source_zip, tmp_path)

    assert source_zip.exists()
    assert any(path.name == "child.txt" for path in extracted)


def test_s13_pipeline_unsupported_xls_records_terminal_outcome(tmp_path):
    from document_processor.parser_factory import ParserFactory
    from document_processor.pipelines.s13_v2_pipeline import S13V2Pipeline

    legacy = tmp_path / "Приложение № 15.xls"
    legacy.write_bytes(b"legacy-binary-xls-fixture")

    pipeline = S13V2Pipeline(
        parser_factory=ParserFactory(),
        downloader=None,
        logger=SimpleNamespace(info=lambda *a, **k: None, warning=lambda *a, **k: None, error=lambda *a, **k: None),
        is_over_memory_limit=lambda: False,
    )

    result = pipeline.process_task(
        queue_id=601,
        procurement_id=202601,
        contract_reg_number="UNSUPPORTED_XLS",
        table_source="reestr_contract_44_fz_awarded",
        files=[legacy],
        match_engine=SimpleNamespace(process_text=lambda *a, **k: []),
    )

    assert result.outcome.value == "SUCCESS"
    assert len(result.files) == 1
    assert result.files[0].file_name == "Приложение № 15.xls"
    assert result.files[0].status == "UNSUPPORTED"
    assert "Unsupported format" in result.files[0].error_message
    assert result.files[0].matches == []


def test_s13_pipeline_parser_exception_records_failed_terminal_outcome(tmp_path):
    from document_processor.pipelines.s13_v2_pipeline import S13V2Pipeline

    doc = tmp_path / "broken.txt"
    doc.write_text("parser should raise", encoding="utf-8")

    class _BrokenParser:
        def parse(self, path):
            raise RuntimeError("boom parser")

    class _Factory:
        def get_parser(self, path):
            return _BrokenParser()

    pipeline = S13V2Pipeline(
        parser_factory=_Factory(),
        downloader=None,
        logger=SimpleNamespace(info=lambda *a, **k: None, warning=lambda *a, **k: None, error=lambda *a, **k: None),
        is_over_memory_limit=lambda: False,
    )

    result = pipeline.process_task(
        queue_id=602,
        procurement_id=202602,
        contract_reg_number="FAILED_PARSE",
        table_source="reestr_contract_44_fz_awarded",
        files=[doc],
        match_engine=SimpleNamespace(process_text=lambda *a, **k: []),
    )

    assert result.outcome.value == "SUCCESS"
    assert result.files[0].status == "FAILED"
    assert "boom parser" in result.files[0].error_message
    assert result.files[0].matches == []


def test_s13_pipeline_archive_input_accounting_docx_xlsx_xls_unsupported(tmp_path):
    import docx

    openpyxl = pytest.importorskip("openpyxl")
    from document_processor.parser_factory import ParserFactory
    from document_processor.pipelines.s13_v2_pipeline import S13V2Pipeline

    source_zip = tmp_path / "source.zip"
    source_zip.write_bytes(b"durable parent archive marker")
    nested = tmp_path / "source" / "docs"
    nested.mkdir(parents=True)

    docx_path = nested / "spec.docx"
    document = docx.Document()
    document.add_paragraph("Документ Word без целевых совпадений.")
    document.save(str(docx_path))

    xlsx_path = nested / "table.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "ТЗ"
    worksheet["A1"] = "Таблица без целевых совпадений"
    workbook.save(xlsx_path)

    legacy_xls = nested / "Приложение № 15.xls"
    legacy_xls.write_bytes(b"legacy xls member")
    blob = nested / "blob.bin"
    blob.write_bytes(b"unsupported binary member")

    pipeline = S13V2Pipeline(
        parser_factory=ParserFactory(),
        downloader=None,
        logger=SimpleNamespace(info=lambda *a, **k: None, warning=lambda *a, **k: None, error=lambda *a, **k: None),
        is_over_memory_limit=lambda: False,
    )

    result = pipeline.process_task(
        queue_id=603,
        procurement_id=202603,
        contract_reg_number="ARCHIVE_MATRIX",
        table_source="reestr_contract_44_fz_awarded",
        files=[docx_path, xlsx_path, legacy_xls, blob],
        match_engine=SimpleNamespace(process_text=lambda *a, **k: []),
    )

    statuses = {file_res.file_name: file_res.status for file_res in result.files}
    assert len(result.files) == 4
    assert statuses["spec.docx"] == "COMPLETED"
    assert statuses["table.xlsx"] == "COMPLETED"
    assert statuses["Приложение № 15.xls"] == "UNSUPPORTED"
    assert statuses["blob.bin"] == "UNSUPPORTED"
    assert sum(1 for file_res in result.files if file_res.status == "COMPLETED") == 2
    assert sum(1 for file_res in result.files if file_res.status == "UNSUPPORTED") == 2
    assert all(file_res.parent_file_name == "source.zip" for file_res in result.files)
    assert sorted(file_res.archive_member_path for file_res in result.files) == [
        "docs/blob.bin",
        "docs/spec.docx",
        "docs/table.xlsx",
        "docs/Приложение № 15.xls",
    ]
