#!/usr/bin/env python3
"""Import approved exact-title rules from an XLSX into crm_procurement_title_rules.

WIP: APPROVED EXACT TITLE RULES V1.

Source of truth: the ``SAFE_EXACT_RULES`` sheet of the classified XLSX.
No Qwen, no re-classification, no TAXONOMY_GAPS handling.

Execution policy (this WIP):
    ASSIGN_SUBCATEGORY -> execution_enabled = TRUE
    KEEP              -> execution_enabled = FALSE
    MOVE_CATEGORY     -> execution_enabled = FALSE
    REMOVE_CATEGORY   -> execution_enabled = FALSE

Usage:
    python scripts/import_approved_title_rules.py --xlsx <file> --dry-run
    python scripts/import_approved_title_rules.py --xlsx <file> --apply
"""
from __future__ import annotations

import argparse
import os
import sys
from collections import Counter

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

from dotenv import load_dotenv  # noqa: E402

load_dotenv(os.path.join(REPO_ROOT, ".env"))

import openpyxl  # noqa: E402

from src.infrastructure.crm_connection import connect_crm  # noqa: E402
from src.services.procurement_title_normalizer import (  # noqa: E402
    NORMALIZATION_VERSION,
    normalize_procurement_title_v1,
)
from src.services.procurement_title_rule_service import (  # noqa: E402
    ACTION_ASSIGN_SUBCATEGORY,
    ACTION_KEEP,
    ACTION_MOVE_CATEGORY,
    ACTION_REMOVE_CATEGORY,
    ALLOWED_ACTIONS,
)

DEFAULT_SHEET = "SAFE_EXACT_RULES"
DEFAULT_APPROVED_SOURCE = "GPT_MANUAL_REVIEW_2026_10_07"

# Only ASSIGN_SUBCATEGORY is executable in this WIP.
EXECUTION_ENABLED = {
    ACTION_ASSIGN_SUBCATEGORY: True,
    ACTION_KEEP: False,
    ACTION_MOVE_CATEGORY: False,
    ACTION_REMOVE_CATEGORY: False,
}

_NEEDED = (
    "category_group_id", "title_group_id", "normalized_title",
    "current_category_code", "target_category_code", "target_subcategory_code",
    "action", "confidence", "rule_scope", "rule_level", "reason",
)


def _cell(row, idx, key):
    v = row[idx[key]] if key in idx else None
    if isinstance(v, str):
        return v.strip()
    return v


def read_rules(xlsx_path, sheet):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"sheet {sheet!r} not found in {xlsx_path}")
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {h: i for i, h in enumerate(header)}
    missing = [c for c in _NEEDED if c not in idx]
    if missing:
        raise SystemExit(f"missing columns in {sheet}: {missing}")
    out = []
    for raw in rows[1:]:
        if raw is None or all(v is None for v in raw):
            continue
        out.append({c: _cell(raw, idx, c) for c in _NEEDED})
    return out


def load_taxonomy(cur):
    cur.execute(
        "SELECT category_code, category_kind, is_active FROM crm_product_categories"
    )
    cats = {r[0]: {"kind": r[1], "active": r[2]} for r in cur.fetchall()}
    cur.execute(
        """
        SELECT c.category_code, s.subcategory_code, s.subcategory_kind, s.is_active
        FROM crm_product_subcategories s
        JOIN crm_product_categories c ON c.id = s.category_id
        """
    )
    subs = {(r[0], r[1]): {"kind": r[2], "active": r[3]} for r in cur.fetchall()}
    return cats, subs


def validate(rule, cats, subs):
    """Return (ok, reason). Rejects are never written."""
    action = (rule.get("action") or "").strip()
    if action not in ALLOWED_ACTIONS:
        return False, f"unknown_action:{action}"
    if not rule.get("current_category_code"):
        return False, "missing_current_category"
    if not rule.get("normalized_title"):
        return False, "missing_normalized_title"

    norm = rule["normalized_title"]
    if normalize_procurement_title_v1(norm) != norm:
        return False, "normalization_not_idempotent"

    try:
        conf = float(rule.get("confidence"))
    except (TypeError, ValueError):
        return False, "bad_confidence"
    if not (0.0 <= conf <= 1.0):
        return False, "confidence_out_of_range"

    if action != ACTION_ASSIGN_SUBCATEGORY:
        return True, ""

    tcat = rule.get("target_category_code")
    tsub = rule.get("target_subcategory_code")
    if not tcat or not tsub:
        return False, "missing_target"
    if tcat != rule["current_category_code"]:
        # Parent-category DML is out of scope in this WIP.
        return False, "target_category_differs_from_current"
    cinfo = cats.get(tcat)
    if not cinfo or not cinfo["active"] or cinfo["kind"] != "PRODUCT":
        return False, "target_category_not_active_product"
    sinfo = subs.get((tcat, tsub))
    if not sinfo or not sinfo["active"] or sinfo["kind"] != "PRODUCT":
        return False, "target_subcategory_not_active_product"
    return True, ""


def build_rows(raw_rules, cats, subs, normalization_version, approved_source):
    rows, rejected = [], []
    for r in raw_rules:
        ok, why = validate(r, cats, subs)
        if not ok:
            rejected.append({"normalized_title": r.get("normalized_title"),
                             "current_category_code": r.get("current_category_code"),
                             "action": r.get("action"), "reason": why})
            continue
        action = r["action"]
        rows.append({
            "normalization_version": normalization_version,
            "normalized_title": r["normalized_title"],
            "current_category_code": r["current_category_code"],
            "action": action,
            "target_category_code": r.get("target_category_code") or None,
            "target_subcategory_code": r.get("target_subcategory_code") or None,
            "confidence": float(r["confidence"]),
            "rule_scope": r.get("rule_scope") or "EXACT_TITLE",
            "rule_level": r.get("rule_level") or None,
            "source_title_group_id": r.get("title_group_id") or None,
            "source_category_group_id": r.get("category_group_id") or None,
            "reason": r.get("reason") or None,
            "approved_source": approved_source,
            "execution_enabled": EXECUTION_ENABLED[action],
        })
    return rows, rejected


UPSERT = """
INSERT INTO crm_procurement_title_rules (
    normalization_version, normalized_title, current_category_code, action,
    target_category_code, target_subcategory_code, confidence, rule_scope,
    rule_level, source_title_group_id, source_category_group_id, reason,
    approved_source, execution_enabled, is_active, created_at, updated_at
) VALUES (
    %(normalization_version)s, %(normalized_title)s, %(current_category_code)s, %(action)s,
    %(target_category_code)s, %(target_subcategory_code)s, %(confidence)s, %(rule_scope)s,
    %(rule_level)s, %(source_title_group_id)s, %(source_category_group_id)s, %(reason)s,
    %(approved_source)s, %(execution_enabled)s, TRUE, NOW(), NOW()
)
ON CONFLICT (normalization_version, normalized_title, current_category_code)
DO UPDATE SET
    action = EXCLUDED.action,
    target_category_code = EXCLUDED.target_category_code,
    target_subcategory_code = EXCLUDED.target_subcategory_code,
    confidence = EXCLUDED.confidence,
    rule_scope = EXCLUDED.rule_scope,
    rule_level = EXCLUDED.rule_level,
    source_title_group_id = EXCLUDED.source_title_group_id,
    source_category_group_id = EXCLUDED.source_category_group_id,
    reason = EXCLUDED.reason,
    approved_source = EXCLUDED.approved_source,
    execution_enabled = EXCLUDED.execution_enabled,
    is_active = TRUE,
    updated_at = NOW()
"""


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--sheet", default=DEFAULT_SHEET)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--normalization-version", default=NORMALIZATION_VERSION)
    ap.add_argument("--approved-source", default=DEFAULT_APPROVED_SOURCE)
    args = ap.parse_args(argv)
    if args.apply == args.dry_run:
        print("choose exactly one of --dry-run / --apply")
        return 2

    raw_rules = read_rules(args.xlsx, args.sheet)
    conn = connect_crm()
    cur = conn.cursor()
    cats, subs = load_taxonomy(cur)
    rows, rejected = build_rows(raw_rules, cats, subs,
                                args.normalization_version, args.approved_source)

    by_action = Counter(r["action"] for r in rows)
    exec_by_action = Counter(r["action"] for r in rows if r["execution_enabled"])

    print(f"RULES_SOURCE={os.path.basename(args.xlsx)}/{args.sheet}")
    print(f"RULES_PARSED={len(raw_rules)}")
    print(f"RULES_IMPORTED={len(rows)}")
    print(f"RULES_BY_ACTION={dict(by_action)}")
    print(f"EXECUTION_ENABLED={dict(exec_by_action)}")
    print(f"IMPORT_REJECTED={len(rejected)}")
    for r in rejected:
        print(f"  REJECT {r['action']} {r['current_category_code']}/{r['normalized_title']}: {r['reason']}")

    if args.apply:
        written = 0
        for r in rows:
            cur.execute(UPSERT, r)
            written += cur.rowcount
        conn.commit()
        print(f"DB_WRITES={written}")
    else:
        conn.rollback()
        print("DB_WRITES=0")
    cur.close()
    conn.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
